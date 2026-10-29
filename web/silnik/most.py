# -*- coding: utf-8 -*-
"""PMT Planer (web) — klej między interfejsem a silnikiem. Odpowiednik `_dane_do_generacji` (nowy_wyglad.py)
i `GeneratorThread.run` (PMT_Delegacje.py) bez Qt. Tu NIE MA rachunku: każdy rachunek to wywołanie funkcji silnika.
Wejście i wyjście: zwykłe słowniki (JSON)."""
import os, json, datetime, math, random
import pmt_silnik as S

# ustawienia z interfejsu nadpisują odczyt pliku ustawień programu (host dostarcza je zamiast okna ustawień)
USTAWIENIA_WEB = {}
_USTAWIENIE_ORYGINAL = S.ustawienie
S.ustawienie = lambda klucz, domyslne="": USTAWIENIA_WEB[klucz] if klucz in USTAWIENIA_WEB else _USTAWIENIE_ORYGINAL(klucz, domyslne)

# ---- „cel wyjazdu” na dokumencie ----
# Silnik (jak program na komputerze) drukuje tu stałą listę sieci — na każdym poleceniu tę samą. Interfejs web losuje dla KAŻDEGO polecenia osobny zestaw:
# zawsze Biedronka i Żabka oraz 3–4 sieci z listy właściciela (silnik/sieci_cel.txt), w kolejności alfabetycznej (po polsku). Losowanie jest powtarzalne
# (zależy od PESEL-u, miesiąca i numeru polecenia), więc ponowne wygenerowanie daje te same sieci. Silnika nie edytujemy (ma być wierną kopią programu),
# dlatego napis podmienia podklasa PDFReport z tej warstwy. Dane pokazowe ani plan wizyt nie mają na to wpływu.
_PREFIKS_CELU = "projekt: Biedronka, Dino"
_CELE_DOKUMENTOW = []
_ALFABET_PL = "aąbcćdeęfghijklłmnńoópqrsśtuvwxyzźż"
_SIECI_STALE = ("Biedronka", "Żabka")
_SIECI_DOMYSLNE = ("Biedronka", "Dino", "Eurocash", "Społem", "Stokrotka", "Żabka")     # stała lista programu — gdy brak pliku z listą


def _klucz_pl(tekst):
    return [(_ALFABET_PL.index(c) if c in _ALFABET_PL else 100 + ord(c)) for c in str(tekst).lower()]


def _wczytaj_liste_sieci():
    """Sieci z pliku sieci_cel.txt (jedna w wierszu, # = komentarz); bez pliku — dotychczasowa lista programu."""
    sciezka = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sieci_cel.txt")
    wynik, widziane = [], set()
    try:
        with open(sciezka, encoding="utf-8-sig") as f:
            for linia in f:
                n = " ".join(linia.replace("\xa0", " ").split())
                if n and not n.startswith("#") and n.lower() not in widziane:
                    widziane.add(n.lower()); wynik.append(n[:1].upper() + n[1:])
    except OSError:
        pass
    return wynik or list(_SIECI_DOMYSLNE)


def _wylosuj_sieci(pesel, rok, mies, nr):
    """Biedronka + Żabka + 3–4 losowe sieci z listy; alfabetycznie. Powtarzalne dla tych samych danych."""
    stale = {s.lower().replace("ż", "z") for s in _SIECI_STALE}
    reszta = [s for s in _wczytaj_liste_sieci() if s.lower().replace("ż", "z") not in stale]
    rng = random.Random("%s|%d|%d|%d" % (pesel, int(rok), int(mies), int(nr)))
    ile = min(len(reszta), rng.choice((3, 4)))
    return sorted(list(_SIECI_STALE) + rng.sample(reszta, ile), key=_klucz_pl)


_MIEJSCA = {}       # numer dokumentu → gdzie na stronach A4 (w mm) jest miejsce na podpis pracownika — z tego aplikacja sprawdza podpis odręczny na skanie


class _PDFReportWeb(S.PDFReport):
    licznik = 0          # numer kolejnego dokumentu w tym przebiegu (zerowany przed generuj_pdfy)

    def __init__(self):
        super().__init__()
        self._nr_dok = _PDFReportWeb.licznik
        _PDFReportWeb.licznik += 1
        _MIEJSCA[self._nr_dok] = {"strony": 1, "miejsca": [], "ostatni_y": 0.0}

    def _zapisz_miejsce(self, w, h, text):
        m = _MIEJSCA.get(self._nr_dok)
        if m is None:
            return
        m["strony"] = max(m["strony"], int(self.page_no()))
        m["ostatni_y"] = max(m["ostatni_y"], float(self.get_y()) + float(h or 0)) if self.page_no() == m["strony"] else m["ostatni_y"]
        if text == "podpis pracownika":
            # komórka z podpisem stoi tuż pod linią podpisu: x, y(linii), szerokość; miejsce na podpis jest nad linią
            m["miejsca"].append({"strona": int(self.page_no()), "x": round(float(self.get_x()), 2), "y": round(float(self.get_y()), 2), "w": round(float(w or 0), 2), "opis": text})

    def cell(self, w=None, h=None, text="", *a, **k):
        if isinstance(text, str):
            self._zapisz_miejsce(w, h, text)
        if isinstance(text, str) and text.startswith(_PREFIKS_CELU) and self._nr_dok < len(_CELE_DOKUMENTOW) and _CELE_DOKUMENTOW[self._nr_dok]:
            text = self._cel_do_rubryki(list(_CELE_DOKUMENTOW[self._nr_dok]), w or 95.0)
        return super().cell(w, h, text, *a, **k)

    def _cel_do_rubryki(self, sieci, szerokosc):
        """„projekt: A, B, C” dopasowane do szerokości rubryki: najpierw mniejsza czcionka, w ostateczności skrócona lista („… i inne”)."""
        limit = szerokosc - 2.0
        tekst = "projekt: " + ", ".join(sieci)
        rozm = 8.0
        while rozm > 5.5 and self.get_string_width(tekst) > limit:
            rozm -= 0.5
            self.set_font_size(rozm)
        while len(sieci) > 1 and self.get_string_width(tekst) > limit:
            sieci.pop()
            tekst = "projekt: " + ", ".join(sieci) + " i inne"
        return tekst


S.PDFReport = _PDFReportWeb

STAWKA_POWYZEJ_900 = 1.15
STAWKA_DO_900 = 0.89


def stawka_dla(silnik_powyzej_900):
    """Jak w programie: silnik_idx 0 (≤ 900 cm³) → 0,89; inaczej 1,15 (PMT_Delegacje.py: stawka = 0.89 if idx == 0 else 1.15)."""
    return STAWKA_POWYZEJ_900 if silnik_powyzej_900 else STAWKA_DO_900


def dni_robocze(rok, miesiac, tryb, wolne=()):
    """dni_robocze_realne() z nowy_wyglad.py + odjęcie dni wolnych użytkownika."""
    S.ustaw_tryb_pracy("wieczory" if str(tryb).lower().startswith("wiecz") else "tydzien")
    return [d for d in S.pobierz_dni_robocze(int(rok), int(miesiac)) if d.day not in set(int(x) for x in wolne)]


def dla_ui_dni(rok, miesiac, tryb, wolne=()):
    ds = dni_robocze(rok, miesiac, tryb, wolne)
    return [d.isoformat() for d in ds]


def przygotuj(p):
    """Kontrola przed generowaniem — te same bramki co `_dane_do_generacji`. Zwraca {ok, powod, ...}."""
    imie = " ".join(str(p.get("imie", "")).split())
    pesel = "".join(str(p.get("pesel", "")).split())
    adres = str(p.get("adres", "")).strip()
    kwota = float(p.get("kwota") or 0)
    rok, mies = int(p["rok"]), int(p["miesiac"])
    stawka = stawka_dla(bool(p.get("silnik_powyzej_900", True)))
    limit_dnia = S.MAX_KWOTA_DNIA
    if not imie:
        return {"ok": False, "powod": "imię i nazwisko"}
    if not S.waliduj_pesel(pesel):
        return {"ok": False, "powod": "PESEL"}
    if not adres:
        return {"ok": False, "powod": "adres"}
    if kwota < S.MIN_KWOTA:
        return {"ok": False, "powod": "min. %s zł" % ("%d" % S.MIN_KWOTA)}
    try:
        adres_d = S.waliduj_adres(adres, p.get("wymus_typ"))
    except ValueError:
        return {"ok": False, "powod": "adres"}
    dni = dni_robocze(rok, mies, p.get("tryb", "tydzien"), p.get("wolne", ()))
    if not dni:
        return {"ok": False, "powod": "brak dni"}
    maks = S.maks_kwota_miesiaca(len(dni), stawka, limit_dnia=limit_dnia)
    if kwota > maks:
        return {"ok": False, "powod": "maks. %s zł" % ("%d" % maks), "maks": maks}
    if not S._menedzer():
        return {"ok": False, "powod": "brak menedzer.txt"}
    woj = S.rozpoznaj_wojewodztwo(adres_d["kod_pocztowy"])
    return {"ok": True, "powod": "", "stawka": stawka, "limit_dnia": limit_dnia, "maks": maks, "liczba_dni": len(dni),
            "woj": woj, "adres_caly": adres_d["adres_caly"], "adres_geo": adres_d.get("adres_geo", adres_d["adres_caly"]),
            "kod_pocztowy": adres_d["kod_pocztowy"], "baza_miasto": adres_d["baza_miasto"], "to_wies": bool(adres_d.get("to_wies")),
            "niejednoznaczne": bool(adres_d.get("niejednoznaczne")), "nazwa_do_pytania": adres_d.get("nazwa_do_pytania", "")}


def _dzien_json(d, stawka):
    """Dzień do interfejsu: etapy z kilometrami (kwota / stawka — tak liczy dokument) i punkty trasy (dla schematu)."""
    etapy = [{"skad": e.skad, "dokad": e.dokad, "wyjazd": e.godz_wyj, "przyjazd": e.godz_przyj,
              "kwota": round(float(e.kwota), 2), "km": round(float(e.kwota) / stawka, 1)} for e in d.etapy]
    surowe = list(getattr(d, "etapy_surowe", None) or [])
    punkty = []
    if surowe:
        punkty.append({"m": surowe[0].skad, "lat": float(surowe[0].skad_lat), "lng": float(surowe[0].skad_lng)})
        for e in surowe:
            punkty.append({"m": e.dokad, "lat": float(e.dokad_lat), "lng": float(e.dokad_lng)})
    return {"data": d.data.isoformat(), "suma": round(float(d.suma), 2), "km": round(sum(x["km"] for x in etapy), 1),
            "dokument": int(getattr(d, "dokument", 0) or 0), "etapy": etapy, "punkty": punkty}


def generuj(p, postep=None):
    """Odpowiednik GeneratorThread.run: współrzędne bazy → generuj_trasy → generuj_pdfy (do folderu w pamięci).
    Zwraca {ok, dni, folder, pliki, zrodlo, niepelna, ...}; błędy jako wyjątki (ValueError z tekstem dla człowieka)."""
    pr = przygotuj(p)
    if not pr["ok"]:
        raise ValueError("Brak danych do generowania: " + pr["powod"])
    cb = postep or (lambda t, v: None)
    rok, mies = int(p["rok"]), int(p["miesiac"])
    stawka = pr["stawka"]
    S.zeruj_zrodlo_odleglosci()
    cb("Pobieranie współrzędnych GPS...", 0.10)
    baza_lat, baza_lng = S.pobierz_coords(pr["adres_geo"], pr["baza_miasto"], pr["woj"])
    imie = " ".join(str(p["imie"]).split()); pesel = "".join(str(p["pesel"]).split())
    prac = S.DanePracownika(imie=imie, pesel=pesel, adres=pr["adres_caly"], stanowisko=p.get("stanowisko") or "KR",
                            kod_pocztowy=pr["kod_pocztowy"], baza_miasto=pr["baza_miasto"], baza_lat=baza_lat,
                            baza_lng=baza_lng, wojewodztwo=pr["woj"])
    cb("Generowanie tras...", 0.40)
    dni = dni_robocze(rok, mies, p.get("tryb", "tydzien"), p.get("wolne", ()))
    wynik = S.generuj_trasy(float(p["kwota"]), pr["baza_miasto"], baza_lat, baza_lng, pr["woj"], dni, pesel, stawka, cb)
    zrodlo = S.stan_zrodla_odleglosci()
    if not wynik:
        raise ValueError("Nie udało się ułożyć ANI JEDNEGO dnia trasy. Najczęstsza przyczyna to adres zamieszkania, "
                         "którego nie da się odnaleźć na mapie — sprawdź ulicę, numer i kod pocztowy. "
                         "Żaden dokument nie został utworzony.")
    cb("Rysowanie dokumentów PDF...", 0.85)
    ms = S.MIESIACE_PL[mies - 1]
    folder = "/dokumenty/Rozliczenie_%s_%s_%dr" % (S.nazwa_do_pliku(prac.imie), ms, rok)
    global _CELE_DOKUMENTOW
    _PDFReportWeb.licznik = 0
    _MIEJSCA.clear()
    _CELE_DOKUMENTOW = [_wylosuj_sieci(pesel, rok, mies, nr) for nr, _ in enumerate(S._podziel_na_dokumenty(sorted(wynik, key=lambda d: d.data)), start=1)]
    S.generuj_pdfy(wynik, prac, mies, rok, folder, stawka, cb, zrodlo=zrodlo)
    pliki = sorted(f for f in os.listdir(folder) if f.lower().endswith(".pdf"))
    delegacje = [f for f in pliki if f.startswith("delegacja_")]
    dokumenty = []
    for nr, f in enumerate(delegacje, start=1):
        dd = [d for d in wynik if int(getattr(d, "dokument", 0) or 0) == nr]
        if dd:
            dokumenty.append({"nr": nr, "plik": f, "od": min(d.data.day for d in dd), "do": max(d.data.day for d in dd),
                              "kw": round(sum(float(d.suma) for d in dd), 2),
                              "sieci": list(_CELE_DOKUMENTOW[nr - 1]) if nr - 1 < len(_CELE_DOKUMENTOW) else []})
    rozl = [f for f in pliki if not f.startswith("delegacja_")]
    miejsca_podpisu = {}
    for idx, f in enumerate(delegacje + rozl):          # kolejność plików = kolejność tworzenia dokumentów (delegacje po numerach, na końcu rozliczenie)
        m = _MIEJSCA.get(idx)
        if m:
            miejsca_podpisu[f] = {"strony": m["strony"], "miejsca": m["miejsca"], "ostatni_y": round(m["ostatni_y"], 1)}
    return {"ok": True, "miejsca_podpisu": miejsca_podpisu, "folder": folder, "pliki": pliki, "dokumenty": dokumenty, "rozliczenie": (rozl[0] if rozl else ""),
            "dni": [_dzien_json(d, stawka) for d in wynik], "stawka": stawka,
            "suma": round(sum(d.suma for d in wynik), 2), "zrodlo": zrodlo,
            "niepelna": bool(getattr(wynik, "kwota_niepelna", False)),
            "za_mala": bool(getattr(wynik, "kwota_za_mala", False)),
            "osiagnieta": float(getattr(wynik, "kwota_osiagnieta", 0.0) or 0.0),
            "docelowa": float(getattr(wynik, "kwota_docelowa", p["kwota"]) or 0.0),
            "dokumentow": S.ile_dokumentow(float(p["kwota"]))}


# ---- wejścia dla JavaScriptu (słownik → słownik) ----
def api_dni(p):
    return {"dni": dla_ui_dni(p["rok"], p["miesiac"], p.get("tryb", "tydzien"), p.get("wolne", ()))}


def api_stale(_p=None):
    """Stałe rachunku z silnika (do pokazania w interfejsie; interfejs niczego nie liczy sam)."""
    return {"limit_dnia": S.MAX_KWOTA_DNIA, "limit_dokumentu": S.MAX_KWOTA_DOKUMENTU, "min_kwota": S.MIN_KWOTA,
            "stawka_powyzej": STAWKA_POWYZEJ_900, "stawka_do": STAWKA_DO_900,
            "miesiace": list(S.MIESIACE_PL), "wersja_silnika": S.SYGNATURA_SILNIKA if hasattr(S, "SYGNATURA_SILNIKA") else ""}


def api_maks(p):
    ds = dni_robocze(p["rok"], p["miesiac"], p.get("tryb", "tydzien"), p.get("wolne", ()))
    st = stawka_dla(bool(p.get("silnik_powyzej_900", True)))
    return {"dni": len(ds), "maks": S.maks_kwota_miesiaca(len(ds), st, limit_dnia=S.MAX_KWOTA_DNIA), "stawka": st,
            "dokumentow": S.ile_dokumentow(float(p.get("kwota") or 0)) if p.get("kwota") else 0}


def api_adres(p):
    """Walidacja adresu silnikiem (waliduj_adres z programu)."""
    try:
        a = S.waliduj_adres(str(p.get("adres", "")).strip(), p.get("wymus_typ"))
    except ValueError as e:
        return {"ok": False, "blad": str(e)}
    return {"ok": True, "adres_caly": a["adres_caly"], "kod_pocztowy": a["kod_pocztowy"], "baza_miasto": a["baza_miasto"],
            "woj": S.rozpoznaj_wojewodztwo(a["kod_pocztowy"]), "to_wies": bool(a.get("to_wies")),
            "niejednoznaczne": bool(a.get("niejednoznaczne")), "nazwa_do_pytania": a.get("nazwa_do_pytania", "")}


# ================= PLAN WIZYT (odpowiednik PlanerWizytThread i _wczytaj_pozycje_z_pliku) =================
import re as _re


def _wiersze_z_pliku(sciezka):
    """Jak `_wczytaj_pozycje_z_pliku`: CSV (separator z próbki) albo XLSX (openpyxl, pierwszy arkusz)."""
    wiersze = []
    if sciezka.lower().endswith(".csv"):
        import csv
        with open(sciezka, newline="", encoding="utf-8-sig") as f:
            probka = f.read(2048); f.seek(0)
            sep = ";" if probka.count(";") > probka.count(",") else ","
            for w in csv.reader(f, delimiter=sep):
                wiersze.append(w)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(sciezka, read_only=True, data_only=True)
        ws = wb.active
        for w in ws.iter_rows(values_only=True):
            wiersze.append(list(w))
        wb.close()
    return wiersze


def plan_wczytaj(p):
    """Lista sklepów z pliku → {ok, pozycje:[{adres,siec,miasto}], kolumny:[role], blad}."""
    try:
        wiersze = _wiersze_z_pliku(p["sciezka"])
    except Exception as e:
        return {"ok": False, "blad": "Nie udało się wczytać pliku: %s" % e}
    wiersze = [w for w in wiersze if any(str(c or "").strip() for c in w)]
    if not wiersze:
        return {"ok": False, "blad": "Plik nie zawiera danych."}
    mapa = S._dopasuj_kolumny(wiersze[0])
    if mapa:
        dane = wiersze[1:]
    else:
        mapa = {"miasto": 0}
        dane = wiersze
    pozycje = []
    for w in dane:
        adres, siec = S._wiersz_na_adres(w, mapa)
        if adres:
            i_m = mapa.get("miasto")
            miasto = str(w[i_m] if (i_m is not None and i_m < len(w) and w[i_m] is not None) else "").strip()
            pozycje.append({"adres": adres, "siec": siec, "miasto": miasto})
    if not pozycje:
        return {"ok": False, "blad": "Nie rozpoznałem kolumn z lokalizacjami. Upewnij się, że plik ma nagłówki (np. Miasto, Ulica)."}
    return {"ok": True, "pozycje": pozycje, "kolumny": sorted(mapa.keys())}


def _baza_planu(adres_bazy):
    """Punkt startowy dnia (jak w PlanerWizytThread.run): cache dokładnego adresu → miasto z bazy offline → sieć."""
    if not adres_bazy:
        return None
    miasto_b = adres_bazy.split(",")[-1].strip() if "," in adres_bazy else adres_bazy
    miasto_b = _re.sub(r"^\d{2}-\d{3}\s*", "", miasto_b).strip()
    lat_b = lng_b = None
    klucz_b = S._klucz_geo(adres_bazy)
    if klucz_b in S._geo_cache:
        lat_b, lng_b = tuple(S._geo_cache[klucz_b])
    else:
        c = S.coords_z_miasta(miasto_b)
        if c:
            lat_b, lng_b = c
        else:
            try:
                lat_b, lng_b = S.pobierz_coords(adres_bazy, miasto_b or adres_bazy, "mazowieckie")
            except Exception:
                lat_b = lng_b = None
    if lat_b is None:
        return None
    return S.PunktWizyty("Baza (start)", adres_bazy, lat=lat_b, lng=lng_b, miasto=miasto_b)


def _godziny_dnia(dzien, baza, tryb):
    """Szacowane godziny dojazdu do kolejnych sklepów: start dnia + przejazdy (km linii × krętość / prędkość) + czas wizyt z ustawień planowania."""
    ust = S.ustawienia_planowania()
    t = (16 * 60 + 30) if tryb == "wieczory" else 8 * 60
    poprz = baza
    wyj = []
    for w in dzien.wizyty:
        if poprz is not None:
            t += S._czas_przejazdu_min(S._dyst_prosty(poprz, w), ust["predkosc"])
        wyj.append("%02d:%02d" % (int(t // 60) % 24, int(t % 60)))
        t += ust["czas_wizyty"]
        poprz = w
    return wyj


def _plan_json(plan, baza, tryb):
    miesiace = []
    for m in plan.get("miesiace", []):
        dni = []
        for d in m["dni"]:
            dni.append({"data": d.data.isoformat(), "km": round(float(d.km), 1), "minuty": round(float(d.minuty)),
                        "godziny": _godziny_dnia(d, baza, tryb),
                        "wizyty": [{"nazwa": w.nazwa, "adres": w.adres, "siec": w.siec, "lat": w.lat, "lng": w.lng, "miasto": w.miasto}
                                   for w in d.wizyty]})
        miesiace.append({"rok": m["rok"], "miesiac": m["miesiac"], "suma_km": round(float(m.get("suma_km", 0)), 1),
                         "suma_wizyt": m.get("suma_wizyt", 0), "dni": dni})
    return {"miesiace": miesiace, "suma_km": round(float(plan.get("suma_km", 0)), 1), "suma_wizyt": plan.get("suma_wizyt", 0),
            "bez_pozycji": [w.nazwa for w in plan.get("bez_pozycji", [])][:50],
            "nierozplanowane": [w.nazwa for w in plan.get("nierozplanowane", [])][:50],
            "baza": ({"adres": baza.adres, "lat": baza.lat, "lng": baza.lng, "miasto": baza.miasto} if baza else None),
            "cykliczny": bool(plan.get("cykliczny")), "tryb": tryb}


def plan_generuj(p, postep=None):
    """Odpowiednik PlanerWizytThread.run: współrzędne z bazy miast → dogeokodowanie nowych miejscowości (raz na miasto)
    → baza → planuj_wizyty / planuj_cyklicznie. Zwraca plan jako słownik (patrz _plan_json)."""
    cb = postep or (lambda t, v: None)
    pozycje = p["pozycje"]; n = len(pozycje)
    punkty = []
    for i, poz in enumerate(pozycje):
        adres = poz.get("adres", ""); siec = poz.get("siec", ""); miasto = poz.get("miasto", "")
        if not miasto:
            miasto = adres.split(",")[-1].strip() if "," in adres else adres
            miasto = _re.sub(r"^\d{2}-\d{3}\s*", "", miasto).strip()
        coords = S.coords_z_miasta(miasto)
        klucz = S._klucz_geo(adres)
        if klucz in S._geo_cache:
            lat, lng = tuple(S._geo_cache[klucz])
        elif coords:
            lat, lng = coords
        else:
            lat = lng = None
        etykieta = "%s — %s" % (siec, miasto) if (siec and miasto) else (siec or miasto or adres)
        punkty.append(S.PunktWizyty(etykieta, adres, siec=siec, lat=lat, lng=lng, miasto=miasto))
        if i % 200 == 0:
            cb("Lokalizowanie z bazy miast (%d/%d)..." % (i + 1, n), (i + 1) / max(n, 1) * 0.5)
    braki = {}
    for pt in punkty:
        if pt.lat is None and pt.miasto:
            braki.setdefault(S._norm_miasto(pt.miasto), []).append(pt)
    if braki:
        lista = list(braki.items())
        for j, (mnorm, pm) in enumerate(lista):
            cb("Nowe miejscowości (%d/%d)..." % (j + 1, len(lista)), 0.5 + (j + 1) / max(len(lista), 1) * 0.3)
            nazwa = pm[0].miasto
            try:
                lat, lng = S.pobierz_coords(nazwa, nazwa, "mazowieckie", zapisz_cache=False)
            except Exception:
                lat = lng = None
            if lat is not None:
                S._zbuduj_indeks_miast()[mnorm] = (lat, lng)
                for pt in pm:
                    pt.lat, pt.lng = lat, lng
        S.zapisz_geo_cache()
    adres_bazy = str(p.get("adres_bazy") or "").strip()
    cb("Lokalizowanie punktu startowego...", 0.82)
    baza = _baza_planu(adres_bazy)
    cb("Układanie tras...", 0.85)
    rok, mies = int(p["rok"]), int(p["miesiac"])
    S.ustaw_tryb_pracy("wieczory" if str(p.get("tryb", "")).startswith("wiecz") else "tydzien")
    offline = not bool(p.get("siec", True))
    # ustawienia planowania dostarcza interfejs (w programie: okno „Ustawienia planowania"): dzień wieczorny to okno 16:30–22:00 = 5,5 h
    USTAWIENIA_WEB["plan_godziny"] = 5.5 if str(p.get("tryb", "")).startswith("wiecz") else 8.0
    od_daty = datetime.date.fromisoformat(p["od_daty"]) if p.get("od_daty") else None
    if p.get("cykliczny"):
        plan = S.planuj_cyklicznie(punkty, rok, mies, cykl_domyslny=int(p.get("cykl_domyslny") or 4),
                                   cykle_sieci=p.get("cykle_sieci") or {}, cykle_sieci_duze={}, cykle_punktow={},
                                   miesiecy=int(p.get("horyzont") or 3), baza=baza, offline=offline, od_daty=od_daty)
    else:
        plan = S.planuj_wizyty(punkty, rok, mies, baza=baza, offline=offline, od_daty=od_daty)
    plan["cykliczny"] = bool(p.get("cykliczny"))
    cb("Gotowe!", 1.0)
    return _plan_json(plan, baza, p.get("tryb", "tydzien"))
